"""Export shared Markdown to HTML and styled Word; render final Word to matching PDF.

Requires pandoc, python-docx, and LibreOffice. Use the Codex bundled document
Python/runtime when available; SOFFICE can select its bundled renderer.
"""
import os
import shutil
import subprocess
import tempfile
from pathlib import Path

from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT=Path(__file__).resolve().parents[1]
DOCS=ROOT/'docs'
NAMES=['midterm-report','submission_guide','database-design']
soffice=os.getenv('SOFFICE') or shutil.which('soffice')
if not soffice:raise SystemExit('Set SOFFICE to a LibreOffice executable before exporting.')

def style_doc(path):
    doc=Document(path)
    for section in doc.sections:
        section.page_width=Inches(8.5);section.page_height=Inches(11)
        section.top_margin=Inches(.8);section.bottom_margin=Inches(.8)
        section.left_margin=Inches(.8);section.right_margin=Inches(.8)
        section.header_distance=Inches(.3);section.footer_distance=Inches(.3)
        footer=section.footer.paragraphs[0];footer.alignment=1
        run=footer.add_run('PayLink | CS 5600 | ');run.font.size=Pt(8)
        fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');footer._p.append(fld)
    normal=doc.styles['Normal'];normal.font.name='Arial';normal.font.size=Pt(11)
    normal.font.color.rgb=RGBColor.from_string('202D38')
    normal.paragraph_format.space_after=Pt(7);normal.paragraph_format.line_spacing=1.12
    for name in ('Title','Subtitle','Heading 1','Heading 2','Heading 3'):
        st=doc.styles[name];st.font.name='Arial';st.font.color.rgb=RGBColor(0,0,0)
        st.font.size=Pt({'Title':24,'Subtitle':12,'Heading 1':16,'Heading 2':13,'Heading 3':11}[name])
        st.paragraph_format.keep_with_next=True
        st.paragraph_format.space_before=Pt(14 if name!='Title' else 0)
        st.paragraph_format.space_after=Pt(7)
        pp=st.element.find(qn('w:pPr'))
        if pp is not None:
            for border in list(pp.findall(qn('w:pBdr'))):pp.remove(border)
    for p in doc.paragraphs:
        p.paragraph_format.widow_control=True
        if p.style.name.startswith(('Heading','Title','Subtitle')):
            for run in p.runs:run.font.color.rgb=RGBColor(0,0,0)
    for table in doc.tables:
        pr=table._tbl.tblPr
        borders=OxmlElement('w:tblBorders')
        for side in ('top','left','bottom','right','insideH','insideV'):
            edge=OxmlElement('w:'+side);edge.set(qn('w:val'),'single');edge.set(qn('w:sz'),'4');edge.set(qn('w:color'),'D9D9D9');borders.append(edge)
        pr.append(borders)
        margins=OxmlElement('w:tblCellMar')
        for side in ('top','left','bottom','right'):
            item=OxmlElement('w:'+side);item.set(qn('w:w'),'90');item.set(qn('w:type'),'dxa');margins.append(item)
        pr.append(margins)
        for i,row in enumerate(table.rows):
            rp=row._tr.get_or_add_trPr()
            no_split=OxmlElement('w:cantSplit');rp.append(no_split)
            if i==0:rp.append(OxmlElement('w:tblHeader'))
            for cell in row.cells:
                cell.vertical_alignment=1
                tcpr=cell._tc.get_or_add_tcPr()
                shading=OxmlElement('w:shd');shading.set(qn('w:fill'),'EAF0FC' if i==0 else ('F7F9FC' if i%2==0 else 'FFFFFF'));tcpr.append(shading)
                for p in cell.paragraphs:
                    p.paragraph_format.space_after=Pt(3);p.paragraph_format.line_spacing=1.06
                    for run in p.runs:
                        run.font.name='Arial';run.font.size=Pt(9.5);run.font.color.rgb=RGBColor(0,0,0)
                        if i==0:run.bold=True
    # Pandoc can write a width wider than the final text area for raster diagrams.
    max_width=Inches(6.9)
    for shape in doc.inline_shapes:
        if shape.width>max_width:
            ratio=max_width/shape.width;shape.width=int(shape.width*ratio);shape.height=int(shape.height*ratio)
    doc.save(path)

for name in NAMES:
    subprocess.run(['pandoc',name+'.md','-s','--embed-resources','--css','report.css','-o',name+'.html'],cwd=DOCS,check=True)
    subprocess.run(['pandoc',name+'.md','-o',name+'.docx'],cwd=DOCS,check=True)
    style_doc(DOCS/(name+'.docx'))
    with tempfile.TemporaryDirectory(prefix='paylink-docs-') as profile:
        env=os.environ.copy();env['TMPDIR']='/private/tmp' if os.uname().sysname=='Darwin' else tempfile.gettempdir()
        subprocess.run([soffice,'-env:UserInstallation='+Path(profile).as_uri(),'--headless','--convert-to','pdf',
                        '--outdir',str(DOCS),str(DOCS/(name+'.docx'))],env=env,check=True,timeout=120)
    if not (DOCS/(name+'.pdf')).is_file():raise RuntimeError('PDF export missing: '+name)
    print('Exported final matching HTML, DOCX and PDF:',name)
